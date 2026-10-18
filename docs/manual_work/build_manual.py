from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "docs" / "manual_work"
SHOTS = WORK / "screenshots"
OUT = ROOT / "docs" / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Postcards Руководство пользователя.docx"

BLACK = "000000"
DARK_BLUE = "243B53"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "6B7280"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_run_font(run, name="Aptos", size=None, bold=None, color=BLACK):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, txt, end])


def add_body(doc, text, bold_lead=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
        r2.italic = italic
    else:
        r = p.add_run(text)
        set_run_font(r)
        r.italic = italic
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_run_font(r)


def add_steps(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.32)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.1
        n = p.add_run(f"{number}.  ")
        set_run_font(n)
        r = p.add_run(item)
        set_run_font(r)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_screenshot(doc, filename, caption, width=6.35, crop_top=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run()
    shape = r.add_picture(str(SHOTS / filename), width=Inches(width))
    if crop_top:
        pic = shape._inline.graphic.graphicData.pic
        src_rect = OxmlElement("a:srcRect")
        src_rect.set("t", str(int(crop_top * 100000)))
        pic.blipFill.insert(1, src_rect)
        shape.height = int(shape.height * (1 - crop_top))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    rr = cap.add_run(caption)
    set_run_font(rr, size=9, color=MID_GRAY)
    rr.italic = True
    return cap


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=10, bold=True, color="FFFFFF")
    for ridx, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if widths:
                cell.width = Inches(widths[i])
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

title = styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(28)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(12)
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size, before, after in (("Heading 1", 19, 18, 8), ("Heading 2", 14, 12, 5), ("Heading 3", 11.5, 9, 3)):
    st = styles[name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[style_name].font.name = "Aptos"
    styles[style_name].font.size = Pt(11)
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(34)
r = p.add_run("Руководство по установке и работе с приложением Postcards")
set_run_font(r, name="Aptos Display", size=28, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Инсталляция для четырёх компьютеров и редактор PDF")
set_run_font(r, size=14, color=MID_GRAY)

add_screenshot(doc, "02-main-scene.png", "Основной экран с восемью открытками", width=5.9)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Версия приложения 2.0.2\nРедакция руководства 18 октября 2026 года")
set_run_font(r, size=10, color=MID_GRAY)

doc.add_page_break()

add_heading(doc, "Назначение руководства", 1)
add_body(doc, "Это руководство предназначено для технической команды, медиаторов и операторов инсталляции Postcards. Оно описывает установку приложения на четыре компьютера, первичную настройку ролей, подключение к координатору, ежедневный запуск, управление открытками, MIDI, показ PDF, редактирование документов и восстановление системы после финального чёрного экрана.")
add_body(doc, "Короткий порядок запуска: сначала включите ПК 1 и дождитесь запуска координатора, затем запустите приложение на ПК 2–ПК 4, проверьте зелёные индикаторы устройств в меню ПК 1 и только после этого запускайте показ.", bold_lead="Короткий порядок запуска:")
add_body(doc, "На отдельных снимках может отображаться тестовый адрес. Рабочий адрес координатора для всех клиентских компьютеров: http://192.168.88.54:8787.", italic=True)

add_heading(doc, "Содержание", 1)
for line in [
    "1  Подготовка компьютеров и сети",
    "2  Установка приложения",
    "3  Настройка ПК 1",
    "4  Настройка ПК 2 ПК 3 и ПК 4",
    "5  Проверка стенда перед показом",
    "6  Работа посетителя с открытками",
    "7  Работа сценария и PDF окон",
    "8  Редактирование документов",
    "9  История версий",
    "10  MIDI и Reaper",
    "11  Завершение показа и перезапуск",
    "12  Обновление файлов и обслуживание",
    "13  Решение типовых проблем",
    "14  Краткая памятка оператора",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(line)
    set_run_font(r)

doc.add_page_break()

add_heading(doc, "1 Подготовка компьютеров и сети", 1)
add_body(doc, "До установки подготовьте четыре компьютера Windows в одной локальной сети. ПК 1 выполняет роль координатора: на нём запускается сервер приложения, хранится общее состояние, принимаются MIDI команды и сохраняются изменения PDF. ПК 2–ПК 4 подключаются к ПК 1 по сети.")
add_heading(doc, "Что должно быть готово", 2)
add_bullets(doc, [
    "одинаковая актуальная сборка Postcards на всех четырёх компьютерах",
    "стабильное проводное сетевое соединение между компьютерами",
    "постоянный IPv4 адрес ПК 1 или зарезервированный адрес в роутере",
    "четыре актуальных PDF файла для кругов показа",
    "звуковые файлы кругов и финального вибро при необходимости",
    "rtpMIDI и проект Reaper на ПК 1 для режима выставки",
])
add_heading(doc, "Рекомендуемый порядок включения", 2)
add_steps(doc, [
    "Включите сетевое оборудование.",
    "Запустите ПК 1 и приложение Postcards. Дождитесь, пока координатор начнёт отвечать.",
    "Запустите приложение на ПК 2, ПК 3 и ПК 4.",
    "На ПК 1 нажмите M и убедитесь, что все четыре роли отмечены зелёным.",
    "Проверьте MIDI и тестовое открытие PDF до допуска посетителей.",
])

add_heading(doc, "2 Установка приложения", 1)
add_body(doc, "Актуальный установщик находится в файле Postcards_2.0.2_x64-setup.exe. Установка выполняется отдельно на каждом из четырёх компьютеров и требует прав администратора Windows.")
add_steps(doc, [
    "Скопируйте установщик на компьютер. Используйте только файл, полученный от ответственного технического специалиста.",
    "Дважды щёлкните Postcards_2.0.2_x64-setup.exe.",
    "Подтвердите системный запрос Windows на внесение изменений. Если Windows показывает предупреждение о неизвестном источнике, сначала проверьте происхождение файла; не отключайте защиту системы.",
    "В мастере установки оставьте предлагаемую папку или выберите согласованную папку установки. Нажмите Install.",
    "Дождитесь завершения. В комплект включён автономный установщик Microsoft WebView2, поэтому первый запуск может занять несколько минут.",
    "Нажмите Finish и откройте Postcards из меню Пуск или ярлыка, если он создан установщиком.",
])
add_body(doc, "Повторите установку на ПК 1, ПК 2, ПК 3 и ПК 4. Не смешивайте разные сборки: интерфейс и сервер должны быть одной версии.", bold_lead="Важно:")

add_heading(doc, "3 Настройка ПК 1", 1)
add_body(doc, "При первом запуске приложение показывает экран настройки этой машины. Настройка сохраняется. Вернуться к ней позже можно из меню по клавише M.")
add_screenshot(doc, "01-setup.png", "Первый запуск и выбор роли компьютера")
add_steps(doc, [
    "В поле Роль выберите PC1.",
    "Оставьте включённым флажок Эта машина координатор.",
    "Убедитесь, что адрес координатора равен http://localhost:8787.",
    "Оставьте включённым Показывать меню администратора по клавише M.",
    "Нажмите Проверить. Ожидаемый ответ: Координатор отвечает.",
    "Нажмите Сохранить и запустить. Приложение перезапустится в полноэкранном режиме.",
])

add_heading(doc, "4 Настройка ПК 2 ПК 3 и ПК 4", 1)
add_body(doc, "Настройте каждый клиентский компьютер отдельно. Рабочий адрес координатора на ПК 1 зафиксирован: http://192.168.88.54:8787.")
add_table(doc, ["Компьютер", "Роль", "Координатор", "Адрес"], [
    ["ПК 2", "PC2", "флажок выключен", "http://192.168.88.54:8787"],
    ["ПК 3", "PC3", "флажок выключен", "http://192.168.88.54:8787"],
    ["ПК 4", "PC4", "флажок выключен", "http://192.168.88.54:8787"],
], widths=[0.75, 0.75, 1.55, 3.75])
add_steps(doc, [
    "Выберите соответствующую роль PC2, PC3 или PC4.",
    "Снимите флажок Эта машина координатор.",
    "В поле адреса координатора введите http://192.168.88.54:8787.",
    "Нажмите Проверить. Продолжайте только после сообщения Координатор отвечает.",
    "Меню администратора по M обычно достаточно оставить только на ПК 1. Включите его на клиенте, если оператору нужна локальная диагностика.",
    "Нажмите Сохранить и запустить.",
])

add_heading(doc, "5 Проверка стенда перед показом", 1)
add_body(doc, "На ПК 1 нажмите M. Меню открывается поверх сцены и закрывается повторным нажатием M или клавишей Esc.")
add_heading(doc, "Что проверить", 2)
add_bullets(doc, [
    "ws имеет состояние online",
    "в строке Devices роли pc1 pc2 pc3 pc4 отмечены зелёным",
    "адрес координатора указан как http://192.168.88.54:8787",
    "звук кликов включён на нужных машинах",
    "на всех компьютерах установлена версия Postcards 2.0.2",
    "PDF окна можно открыть кнопками open pc1 open pc2 open pc3 open pc4",
    "после проверки окна закрываются кнопками close pc1 close pc2 close pc3 close pc4",
])
add_body(doc, "Если клиент не появляется зелёным, проверьте адрес координатора на клиенте, сетевой кабель, профиль сети Windows и правило файрвола для порта 8787.")

add_heading(doc, "6 Работа посетителя с открытками", 1)
add_screenshot(doc, "02-main-scene.png", "Экран посетителя")
add_body(doc, "На каждом экране расположены восемь открыток. Посетитель нажимает открытку, она переворачивается, и начинается воспроизведение связанного звукового фрагмента. Новый клик во время воспроизведения текущего фрагмента не засчитывается и не прерывает звук.")
add_body(doc, "Клик записывается после окончания фрагмента и с учётом параметра clickGateMs. Когда один из компьютеров набирает 17 засчитанных кликов, приложение запускает сценарий. В режиме выставки ПК 1 отправляет в Reaper MIDI ноту 72.")
add_heading(doc, "Переворот всех открыток", 2)
add_body(doc, "Оператор может одновременно перевернуть все восемь открыток на всех четырёх компьютерах. Для этого нажмите F на клавиатуре или используйте кнопку Перевернуть все открытки без кликов в меню ПК 1. Повторное нажатие возвращает открытки обратно.")
add_body(doc, "Эта служебная операция не увеличивает счётчики кликов, не запускает звук и не приближает порог 17 кликов.", bold_lead="Важно:")

add_heading(doc, "Горячие клавиши", 2)
add_table(doc, ["Клавиша", "Действие", "Где работает"], [
    ["M", "открыть или закрыть меню администратора", "главное окно при включённом меню"],
    ["F", "перевернуть все открытки без учёта кликов", "главное окно и PDF окно"],
    ["P", "свернуть или вернуть PDF окно этой машины", "установленная версия Tauri"],
    ["L", "открыть или закрыть журнал", "главное окно и PDF окно"],
    ["Esc", "закрыть меню или журнал", "соответствующее открытое окно"],
    ["Delete", "удалить выбранный объект", "режим редактирования PDF"],
], widths=[0.75, 3.65, 2.4])
add_body(doc, "Горячие клавиши M и F привязаны к физическим клавишам, поэтому работают и при русской раскладке.")

add_heading(doc, "7 Работа сценария и PDF окон", 1)
add_body(doc, "Сценарий состоит из четырёх кругов. В каждом круге приложение открывает соответствующий файл pdf1 pdf2 pdf3 или pdf4. В режиме выставки открытие и закрытие окон ПК 1–ПК 4 управляется MIDI командами из Reaper. В режиме ноутбуков тот же порядок выполняется таймерами без MIDI.")
add_body(doc, "После завершения круга PDF остаются доступными на осевших компьютерах как вкладки. Новый круг добавляет следующую вкладку. Активный документ выбирается автоматически по роли и этапу сценария.")
add_screenshot(doc, "06-pdf-view.png", "PDF окно в режиме просмотра")
add_heading(doc, "Ручное управление", 2)
add_bullets(doc, [
    "Open PDF local открывает PDF на ПК 1 для проверки",
    "Close PDF local закрывает локальное PDF окно",
    "Force open all принудительно открывает окна на всех ролях",
    "P сворачивает или возвращает PDF окно на текущей машине",
    "TEST RUN выполняет короткий автоматический прогон только на подключённых машинах",
    "Stop test останавливает тестовый прогон",
])
add_body(doc, "Ручные команды open и close переводят сценарий в служебный режим. Перед реальным показом после такой проверки выполните Hard reset.")

add_heading(doc, "8 Редактирование документов", 1)
add_body(doc, "Основная вёрстка четырёх документов выполняется в Postcards Studio по адресу https://mess-studio.ru. Локальный редактор в десктопном приложении остаётся для быстрых исправлений поверх уже установленного PDF.")
add_heading(doc, "Веб редактор Postcards Studio", 2)
add_steps(doc, [
    "Откройте https://mess-studio.ru и войдите под личной учётной записью.",
    "Выберите один из четырёх документов. Рабочие правки сохраняются автоматически и синхронизируются между участниками.",
    "Вставляйте основной текст в пустые поля. Красная строка добавляется автоматически при вставке и после Enter.",
    "Для горизонтальных и вертикальных линий можно снять фиксацию, изменить положение, длину, толщину и вариант пунктира, затем нажать «Зафиксировать и синхронизировать».",
    "Пользователь publisher или admin сохраняет общий шаблон. Геометрия применяется ко всем 14 разворотам и четырём документам, а тексты и изображения каждого документа сохраняются отдельно.",
    "Перед публикацией скачайте PDF и проверьте переносы, красные строки, изображения и все 28 страниц. Предпросмотр использует те же метрики текста, что и PDF.",
    "После проверки нажмите «Опубликовать 4 PDF».",
])
add_heading(doc, "Локальный редактор на ПК 1", 2)
add_steps(doc, [
    "Откройте нужный PDF автоматически по сценарию или вручную из меню.",
    "Нажмите Изменить в верхней панели PDF окна.",
    "Выполните правки. Через 0,9 секунды после последнего изменения приложение автоматически создаёт новую версию.",
    "Дождитесь надписи Сохранено с датой и временем.",
    "Нажмите Готово, чтобы выйти из режима редактирования.",
])
add_screenshot(doc, "07-pdf-editor.png", "Редактор PDF после добавления текстового блока")
add_heading(doc, "Текст", 2)
add_steps(doc, [
    "Нажмите Текст. На текущей странице появится блок Новый текст.",
    "Щёлкните внутри блока и введите текст.",
    "Перемещайте блок за синюю ручку в левом верхнем углу.",
    "Меняйте размер за синюю ручку в правом нижнем углу.",
    "В панели свойств настройте размер, цвет, фон, жирное начертание и выравнивание.",
])
add_heading(doc, "Замена старого текста", 2)
add_body(doc, "Редактор не меняет векторное содержимое исходного PDF. Чтобы заменить уже существующий текст, нажмите Скрыть фрагмент, разместите белую маску поверх старого текста, затем добавьте новый текстовый блок. Исходный PDF остаётся неизменным, а правки хранятся отдельным слоем.")
add_heading(doc, "Изображения", 2)
add_steps(doc, [
    "Нажмите Изображение и выберите PNG JPG WEBP или GIF размером не более 25 МБ.",
    "Переместите изображение за верхнюю левую ручку и измените размер за нижнюю правую.",
    "Выберите Вписать целиком, чтобы показать всё изображение, или Заполнить область, чтобы заполнить рамку с возможной обрезкой.",
    "Используйте кнопки Слой вниз и Слой вверх для порядка наложения.",
    "Чтобы удалить объект, выделите его и нажмите Удалить или Delete.",
])
add_heading(doc, "Одновременная работа", 2)
add_body(doc, "Не редактируйте один и тот же документ одновременно на двух компьютерах. Если другая машина уже сохранила новую версию, приложение блокирует ваше сохранение и показывает сообщение о конфликте. Завершите работу на другой машине, затем нажмите Обновить и повторите свои правки.")

add_heading(doc, "9 История версий", 1)
add_body(doc, "Каждое автоматическое сохранение увеличивает номер версии и записывает дату, время и роль компьютера, с которого выполнена правка. История версий хранится локально на ПК 1.")
add_heading(doc, "Восстановление версии", 2)
add_steps(doc, [
    "Откройте Историю.",
    "Найдите версию по номеру, дате и роли автора.",
    "Нажмите нужную версию. Она загрузится в редактор как текущая правка.",
    "Дождитесь нового автоматического сохранения. Восстановленная версия получит новый номер и не удалит прежнюю историю.",
])

add_heading(doc, "10 MIDI и Reaper", 1)
add_screenshot(doc, "05-midi.png", "Состояние MIDI порта и монитор входящих сообщений", width=5.8, crop_top=0.08)
add_body(doc, "В режиме Выставка MIDI обслуживается сервером на ПК 1. Рекомендуемые параметры стенда: вход PC-10, выход PC-10, фильтрация по каналу 2. Состояние Порт открыто означает, что приложение подключилось к выбранному MIDI входу.")
add_table(doc, ["Нота", "Действие по умолчанию"], [
    ["60", "завершить сценарий после финального удержания или чёрного экрана"],
    ["61", "открыть PDF на ПК 1"],
    ["62", "закрыть PDF на ПК 1"],
    ["63", "открыть PDF на ПК 2"],
    ["64", "закрыть PDF на ПК 2"],
    ["65", "открыть PDF на ПК 3"],
    ["66", "закрыть PDF на ПК 3"],
    ["67", "открыть PDF на ПК 4"],
    ["68", "закрыть PDF на ПК 4"],
    ["69", "немедленно перейти к финальному чёрному экрану"],
    ["72 выход", "сигнал Reaper после достижения порога 17 кликов"],
], widths=[1.2, 5.6])
add_heading(doc, "Проверка без MIDI оборудования", 2)
add_body(doc, "В блоке Имитация входящей ноты укажите ноту и канал, затем нажмите Отправить. Команда проходит через тот же фильтр, дедупликацию и маппинг, что и реальное MIDI сообщение. Для проверки открытия и закрытия удобнее использовать кнопки тест рядом со строками маппинга.")
add_body(doc, "Имитация изменяет реальное состояние сценария. После проверки выполните Hard reset.", bold_lead="Важно:")
add_heading(doc, "Если ноты видны но действие не выполняется", 2)
add_bullets(doc, [
    "проверьте, что включён режим Выставка MIDI Reaper",
    "проверьте выбранный input и состояние Порт открыт",
    "проверьте канал 2 и включённый фильтр",
    "сверьте ноту с таблицей маппинга",
    "учтите дедупликацию одинаковых note on в течение 180 миллисекунд",
    "нажмите L и найдите событие filtered unmapped deduped или dispatch",
])

add_heading(doc, "11 Завершение показа и перезапуск", 1)
add_body(doc, "В финальной фазе приложение закрывает PDF окна и показывает чёрный экран. Чёрный экран расположен ниже служебного меню, поэтому оператор на ПК 1 по прежнему может нажать M.")
add_heading(doc, "Как начать следующий показ", 2)
add_steps(doc, [
    "На ПК 1 нажмите M. Меню откроется поверх чёрного экрана.",
    "Нажмите Hard reset. Команда сбросит сценарий, счётчики кликов, блокировки и положение открыток.",
    "Закройте меню клавишей Esc или M.",
    "Проверьте, что снова видна сцена с открытками и ws остаётся online.",
])
add_body(doc, "Reset scenario завершает только текущий сценарий и предназначен для служебного восстановления. Для нового полного прогона используйте Hard reset.", bold_lead="Разница сбросов:")

add_heading(doc, "12 Обновление файлов и обслуживание", 1)
add_screenshot(doc, "04-media-mode.png", "Замена PDF и звуков и выбор режима инсталляции")
add_heading(doc, "Синхронизация PDF с NAS", 2)
add_steps(doc, [
    "Остановите сценарий и убедитесь, что PDF окна закрыты.",
    "Проверьте доступность диска Z: и файла current.json в папке Z:\\Yandex.Disk\\Выставка Анастасии Чермантеевой\\Postcards Studio\\Релизы.",
    "На ПК 1 откройте меню M и найдите раздел Документы и звуки.",
    "Нажмите «Синхронизировать все документы». Приложение проверит SHA-256 четырёх файлов, сохранит резервную копию прежнего комплекта и атомарно установит pdf1–pdf4.",
    "Дождитесь сообщения с идентификатором установленного релиза.",
    "Откройте все четыре PDF вручную и проверьте первую и последнюю страницы, шрифты и изображения.",
])
add_body(doc, "Если NAS подключён под другой буквой или путём, задайте на ПК 1 переменную среды POSTCARDS_RELEASES_DIR. Публикация на сайте не заменяет локальные PDF до нажатия кнопки синхронизации.", bold_lead="Важно:")
add_heading(doc, "Ручная замена одного PDF", 2)
add_body(doc, "Кнопка Выбрать рядом с отдельным документом остаётся резервным способом. Используйте её только для точечной проверки или аварийной замены; штатное обновление выполняйте кнопкой «Синхронизировать все документы».")
add_body(doc, "Если во время подмены файл занят выдачей клиенту, сервер может вернуть ошибку 409. Закройте PDF окна, подождите несколько секунд и повторите.")
add_heading(doc, "Замена звуков", 2)
add_body(doc, "В том же разделе можно выбрать звук для каждого круга и финального вибро. Кнопка с треугольником воспроизводит тест, а кнопка с крестиком возвращает встроенный файл. Меняйте файлы в состоянии idle.")
add_heading(doc, "Режим работы", 2)
add_bullets(doc, [
    "Выставка MIDI Reaper использует MIDI последовательность с ПК 1",
    "Ноутбуки выполняют первый круг и остальные этапы по таймерам без MIDI",
])

add_heading(doc, "13 Решение типовых проблем", 1)
add_table(doc, ["Симптом", "Проверка и действие"], [
    ["PDF не открывается", "На ПК 1 проверьте Devices, затем откройте роль вручную. Проверьте наличие pdf1–pdf4 и при необходимости повторите «Синхронизировать все документы»."],
    ["MIDI ноты приходят но окна не меняются", "Проверьте режим Выставка, канал, маппинг и фазу сценария. Откройте журнал L и найдите причину filtered, unmapped, deduped или ignored."],
    ["Клиент offline", "Проверьте адрес http://192.168.88.54:8787, кабель и файрвол. На клиенте вернитесь к настройке роли и нажмите Проверить."],
    ["Нет звука кликов", "Включите Звук кликов, выберите аудиоустройство и нажмите тест. На Windows проверьте микшер громкости для Postcards."],
    ["Изменения PDF не сохраняются", "Дождитесь статуса Сохранено. При конфликте завершите редактирование на другой машине и нажмите Обновить."],
    ["После финала только чёрный экран", "Это штатное состояние. Нажмите M, затем Hard reset."],
    ["После теста сценарий ведёт себя неправильно", "Выполните Hard reset и убедитесь, что тестовый режим остановлен."],
], widths=[2.05, 4.75])
add_heading(doc, "Журнал диагностики", 2)
add_body(doc, "Клавиша L открывает журнал клиента и решений координатора. При обращении к разработчику сохраните время ошибки, роль компьютера, фазу сценария, входящую MIDI ноту и несколько строк журнала до и после события.")

add_heading(doc, "14 Краткая памятка оператора", 1)
add_heading(doc, "Перед показом", 2)
add_bullets(doc, [
    "ПК 1 запущен первым",
    "pc1 pc2 pc3 pc4 зелёные в Devices",
    "все четыре PDF открываются вручную",
    "установлен последний опубликованный релиз с NAS",
    "звук и вибро проверены",
    "MIDI порт открыт и тестовые ноты выполняют команды",
    "выполнен Hard reset",
])
add_heading(doc, "Во время показа", 2)
add_bullets(doc, [
    "не меняйте сеть и роли компьютеров",
    "не редактируйте один PDF одновременно на двух машинах",
    "для массового переворота используйте F",
    "для временного сворачивания PDF используйте P",
    "для диагностики используйте L",
])
add_heading(doc, "После показа", 2)
add_bullets(doc, [
    "на чёрном экране нажмите M",
    "для следующего полного прогона нажмите Hard reset",
    "проверьте статус сохранения PDF и наличие новых версий в истории",
])

# Footer with document label and page number.
for sec in doc.sections:
    footer = sec.footer
    footer.is_linked_to_previous = False
    table = footer.add_table(rows=1, cols=2, width=Inches(7.0))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(5.8)
    right.width = Inches(1.2)
    for cell in (left, right):
        set_cell_margins(cell, top=40, bottom=20, start=0, end=0)
    lp = left.paragraphs[0]
    lr = lp.add_run("Postcards  Руководство пользователя")
    set_run_font(lr, size=8.5, color=MID_GRAY)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = rp.add_run("Страница ")
    set_run_font(rr, size=8.5, color=MID_GRAY)
    add_field(rp, "PAGE")

doc.core_properties.title = "Руководство по установке и работе с приложением Postcards"
doc.core_properties.subject = "Установка, настройка, работа, MIDI и редактор PDF"
doc.core_properties.author = "Команда проекта Postcards"
doc.core_properties.keywords = "Postcards, инструкция, установка, MIDI, PDF"
doc.save(DOCX)
print(DOCX)
