from __future__ import annotations

import json
import hmac
import re
import shutil
import uuid
from collections import Counter
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from docx import Document as WordDocument
from docx.oxml.ns import qn
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from starlette.background import BackgroundTask
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .config import settings
from .db import (Document, DocumentRevision, DocumentTemplate, Project, Release,
                 SessionLocal, User, init_db, session_scope)
from .pdf_renderer import publish_release, render_document
from .schemas import (CollaborationAuthBody, CollaborationSnapshotBody, CreateUserBody,
                      LoginBody, RestoreBody, SaveDocumentBody, SaveTemplateBody)
from .security import (current_user, decode_token, hash_password, make_token, publisher,
                       sync_access, verify_password)


cfg = settings()
assets_root = cfg.storage_dir / "assets"
releases_root = cfg.storage_dir / "releases"


def bootstrap_database() -> None:
    cfg.validate_runtime()
    init_db()
    with SessionLocal() as db:
        email = cfg.admin_email.strip().lower()
        admin = db.scalar(select(User).where(User.email == email))
        if not admin:
            db.add(User(email=email, name="Администратор", role="admin",
                        password_hash=hash_password(cfg.admin_password)))
        project = db.scalar(select(Project).order_by(Project.id))
        if not project:
            project = Project(name="Postcards")
            db.add(project)
            db.flush()
            for number in range(1, 5):
                db.add(Document(project_id=project.id, slot=f"pdf{number}",
                                title=f"Текст для ПК{number}", content_json='{"pages":{}}'))
        db.commit()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    bootstrap_database()
    yield


app = FastAPI(title="Postcards Studio API", version="0.1.0", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", cfg.public_url],
                   allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

assets_root.mkdir(parents=True, exist_ok=True)
releases_root.mkdir(parents=True, exist_ok=True)
app.mount("/assets", StaticFiles(directory=str(assets_root)), name="assets")


def now() -> datetime:
    return datetime.now(timezone.utc)


def iso(value: datetime | None) -> str | None:
    if not value:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.isoformat(timespec="milliseconds")


def user_view(user: User) -> dict:
    return {"id": user.id, "email": user.email, "name": user.name, "role": user.role}


def document_view(document: Document) -> dict:
    try:
        content = json.loads(document.content_json)
    except Exception:
        content = {"pages": {}}
    return {"id": document.id, "projectId": document.project_id, "slot": document.slot,
            "title": document.title, "revision": document.revision,
            "updatedAt": iso(document.updated_at), "updatedBy": document.updated_by,
            "pages": content.get("pages") or {}}


def clean_pages(raw_pages: dict) -> dict:
    pages = {}
    for raw_number, raw_items in (raw_pages or {}).items():
        try:
            page_number = int(raw_number)
        except (TypeError, ValueError):
            continue
        if not 1 <= page_number <= 28 or not isinstance(raw_items, list):
            continue
        cleaned = []
        for raw in raw_items[:500]:
            if not isinstance(raw, dict) or raw.get("type") not in {"text", "image", "cover", "shape", "page-background"}:
                continue
            item = {"id": re.sub(r"[^a-zA-Z0-9_-]", "", str(raw.get("id") or uuid.uuid4().hex))[:80],
                    "type": raw["type"]}
            for key, default, lower, upper in (("x", 0.1, 0, 1), ("y", 0.1, 0, 1),
                                                ("w", 0.3, 0.01, 1), ("h", 0.1, 0.01, 1)):
                try:
                    value = float(raw.get(key, default))
                except (TypeError, ValueError):
                    value = default
                item[key] = max(lower, min(upper, value))
            if item["type"] == "text":
                item.update(text=str(raw.get("text") or "")[:30000],
                            fontSize=max(6, min(240, float(raw.get("fontSize") or 28))),
                            fontFamily=raw.get("fontFamily") if raw.get("fontFamily") in {"Arial", "Georgia", "Courier New", "Times New Roman"} else "Times New Roman",
                            color=str(raw.get("color") or "#111111")[:32],
                            background=str(raw.get("background") or "transparent")[:32],
                            align=raw.get("align") if raw.get("align") in {"left", "center", "right", "justify"} else "left",
                            bold=bool(raw.get("bold")), italic=bool(raw.get("italic")))
            elif item["type"] == "image":
                source = str(raw.get("src") or "")
                if source and not source.startswith("/assets/"):
                    continue
                if not source and raw.get("role") != "scan":
                    continue
                item.update(src=source[:600], fit="cover" if raw.get("fit") == "cover" else "contain")
            elif item["type"] in {"cover", "page-background"}:
                item["color"] = str(raw.get("color") or "#ffffff")[:32]
            else:
                item.update(
                    shape=raw.get("shape") if raw.get("shape") in {"rectangle", "ellipse", "line"} else "rectangle",
                    fill=str(raw.get("fill") or "transparent")[:32],
                    stroke=str(raw.get("stroke") or "#111111")[:32],
                    strokeWidth=max(1, min(20, float(raw.get("strokeWidth") or 2))),
                    strokeStyle=raw.get("strokeStyle") if raw.get("strokeStyle") in {"solid", "dashed", "dotted"} else "solid",
                )
            item["locked"] = bool(raw.get("locked"))
            if raw.get("role") in {"scan", "caption", "body", "outline", "divider"}:
                item["role"] = raw["role"]
            if "placeholder" in raw:
                item["placeholder"] = str(raw.get("placeholder") or "")[:300]
            if "sourceLabel" in raw:
                item["sourceLabel"] = str(raw.get("sourceLabel") or "")[:500]
            for key in ("sequence", "cardIndex"):
                try:
                    item[key] = max(1, min(9999, int(raw[key])))
                except (KeyError, TypeError, ValueError):
                    pass
            cleaned.append(item)
        if cleaned:
            pages[str(page_number)] = cleaned
    return pages


class ProjectHub:
    def __init__(self):
        self.connections: dict[int, set[WebSocket]] = {}

    async def connect(self, project_id: int, socket: WebSocket):
        await socket.accept()
        self.connections.setdefault(project_id, set()).add(socket)

    def disconnect(self, project_id: int, socket: WebSocket):
        self.connections.get(project_id, set()).discard(socket)

    async def publish(self, project_id: int, payload: dict):
        dead = []
        for socket in self.connections.get(project_id, set()):
            try:
                await socket.send_json(payload)
            except Exception:
                dead.append(socket)
        for socket in dead:
            self.disconnect(project_id, socket)


hub = ProjectHub()


@app.get("/api/health")
def health():
    return {"ok": True, "service": "postcards-studio", "version": "0.1.0"}


@app.post("/api/auth/login")
def login(body: LoginBody, db: Session = Depends(session_scope)):
    user = db.scalar(select(User).where(User.email == body.email.strip().lower()))
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(401, "Неверная почта или пароль")
    return {"token": make_token(user), "user": user_view(user)}


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return user_view(user)


@app.post("/api/users")
def create_user(body: CreateUserBody, admin: User = Depends(current_user),
                db: Session = Depends(session_scope)):
    if admin.role != "admin":
        raise HTTPException(403, "Только администратор создаёт пользователей")
    if body.role not in {"editor", "publisher", "admin"}:
        raise HTTPException(400, "Неизвестная роль")
    email = body.email.strip().lower()
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(409, "Пользователь уже существует")
    value = User(email=email, name=body.name.strip() or email, role=body.role,
                 password_hash=hash_password(body.password))
    db.add(value)
    db.commit()
    return user_view(value)


@app.get("/api/users")
def users(admin: User = Depends(current_user), db: Session = Depends(session_scope)):
    if admin.role != "admin":
        raise HTTPException(403, "Только администратор видит пользователей")
    return [user_view(value) for value in db.scalars(select(User).order_by(User.name, User.email)).all()]


@app.delete("/api/users/{user_id}")
def delete_user(user_id: int, admin: User = Depends(current_user),
                db: Session = Depends(session_scope)):
    if admin.role != "admin":
        raise HTTPException(403, "Только администратор удаляет пользователей")
    value = db.get(User, user_id)
    if not value:
        raise HTTPException(404, "Пользователь не найден")
    if value.id == admin.id:
        raise HTTPException(409, "Нельзя удалить собственную учётную запись")
    if value.email == cfg.admin_email.strip().lower():
        raise HTTPException(409, "Нельзя удалить основного администратора")
    if value.role == "admin":
        admins = db.scalar(select(func.count(User.id)).where(User.role == "admin")) or 0
        if admins <= 1:
            raise HTTPException(409, "Нельзя удалить последнего администратора")
    db.delete(value)
    db.commit()
    return {"deleted": user_id}


@app.get("/api/projects")
def projects(user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    return [{"id": value.id, "name": value.name, "createdAt": iso(value.created_at)}
            for value in db.scalars(select(Project).order_by(Project.id)).all()]


@app.get("/api/projects/{project_id}")
def project(project_id: int, user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    value = db.get(Project, project_id)
    if not value:
        raise HTTPException(404, "Проект не найден")
    documents = db.scalars(select(Document).where(Document.project_id == project_id).order_by(Document.slot)).all()
    return {"id": value.id, "name": value.name,
            "documents": [{key: doc_value[key] for key in ("id", "slot", "title", "revision", "updatedAt", "updatedBy")}
                          for item in documents for doc_value in [document_view(item)]]}


def find_document(db: Session, project_id: int, slot: str) -> Document:
    document = db.scalar(select(Document).where(Document.project_id == project_id, Document.slot == slot))
    if not document:
        raise HTTPException(404, "Документ не найден")
    return document


def parse_collaboration_name(value: str) -> tuple[int, str]:
    match = re.fullmatch(r"project-(\d+)-(pdf[1-4])", value)
    if not match:
        raise HTTPException(400, "Некорректное имя совместного документа")
    return int(match.group(1)), match.group(2)


def require_collaboration_service(x_collaboration_secret: str = Header(default="")) -> None:
    if not x_collaboration_secret or not hmac.compare_digest(x_collaboration_secret, cfg.collaboration_secret):
        raise HTTPException(401, "Недействительный служебный токен")


def history_checkpoint_due(document: Document, db: Session) -> bool:
    latest = db.scalar(select(DocumentRevision).where(DocumentRevision.document_id == document.id)
                       .order_by(DocumentRevision.created_at.desc()).limit(1))
    if not latest:
        return True
    created_at = latest.created_at
    if created_at.tzinfo is None:
        created_at = created_at.replace(tzinfo=timezone.utc)
    return now() - created_at >= timedelta(minutes=max(1, cfg.history_interval_minutes))


def add_history_checkpoint(document: Document, updated_by: str, db: Session) -> None:
    existing = db.scalar(select(DocumentRevision.id).where(
        DocumentRevision.document_id == document.id,
        DocumentRevision.revision == document.revision,
    ))
    if not existing:
        db.add(DocumentRevision(document_id=document.id, revision=document.revision,
                                content_json=document.content_json, created_by=updated_by))


async def store_collaboration_snapshot(document: Document, pages_value: dict, updated_by: str,
                                       db: Session, checkpoint: bool = False) -> dict:
    pages = clean_pages(pages_value)
    payload = json.dumps({"pages": pages}, ensure_ascii=False, separators=(",", ":"))
    if len(payload.encode("utf-8")) > 4 * 1024 * 1024:
        raise HTTPException(413, "Документ слишком большой")
    if payload == document.content_json:
        if checkpoint:
            add_history_checkpoint(document, updated_by, db)
            db.commit()
        return document_view(document)
    make_checkpoint = checkpoint or history_checkpoint_due(document, db)
    document.revision += 1
    document.content_json = payload
    document.updated_at = now()
    document.updated_by = updated_by
    if make_checkpoint:
        add_history_checkpoint(document, updated_by, db)
    db.commit()
    await hub.publish(document.project_id, {"type": "document.saved", "slot": document.slot,
                                            "revision": document.revision, "updatedBy": updated_by})
    return document_view(document)


@app.post("/api/internal/collaboration/auth")
def collaboration_auth(body: CollaborationAuthBody,
                       access: None = Depends(require_collaboration_service),
                       db: Session = Depends(session_scope)):
    del access
    payload = decode_token(body.token)
    user = db.get(User, int(payload["sub"]))
    if not user:
        raise HTTPException(401, "Пользователь не найден")
    project_id, slot = parse_collaboration_name(body.documentName)
    find_document(db, project_id, slot)
    return {"user": user_view(user), "projectId": project_id, "slot": slot}


@app.get("/api/internal/collaboration/documents/{document_name}")
def collaboration_document(document_name: str,
                           access: None = Depends(require_collaboration_service),
                           db: Session = Depends(session_scope)):
    del access
    project_id, slot = parse_collaboration_name(document_name)
    return document_view(find_document(db, project_id, slot))


@app.post("/api/internal/collaboration/documents/{document_name}")
async def collaboration_store(document_name: str, body: CollaborationSnapshotBody,
                              access: None = Depends(require_collaboration_service),
                              db: Session = Depends(session_scope)):
    del access
    project_id, slot = parse_collaboration_name(document_name)
    document = find_document(db, project_id, slot)
    return await store_collaboration_snapshot(document, body.pages, body.updatedBy or "collaboration", db,
                                              checkpoint=body.checkpoint)


@app.post("/api/projects/{project_id}/documents/{slot}/collaboration-snapshot")
async def collaboration_client_snapshot(project_id: int, slot: str, body: CollaborationSnapshotBody,
                                        user: User = Depends(current_user),
                                        db: Session = Depends(session_scope)):
    document = find_document(db, project_id, slot)
    return await store_collaboration_snapshot(document, body.pages, user.email, db,
                                              checkpoint=body.checkpoint)


@app.get("/api/projects/{project_id}/documents/{slot}")
def get_document(project_id: int, slot: str, user: User = Depends(current_user),
                 db: Session = Depends(session_scope)):
    del user
    return document_view(find_document(db, project_id, slot))


@app.put("/api/projects/{project_id}/documents/{slot}")
async def save_document(project_id: int, slot: str, body: SaveDocumentBody,
                        user: User = Depends(current_user), db: Session = Depends(session_scope)):
    document = find_document(db, project_id, slot)
    if body.baseRevision != document.revision:
        raise HTTPException(409, detail={"message": "Документ изменён другим медиатором",
                                         "current": document_view(document)})
    pages = clean_pages(body.pages)
    payload = json.dumps({"pages": pages}, ensure_ascii=False, separators=(",", ":"))
    if len(payload.encode("utf-8")) > 4 * 1024 * 1024:
        raise HTTPException(413, "Документ слишком большой")
    document.revision += 1
    document.content_json = payload
    document.updated_at = now()
    document.updated_by = user.email
    db.add(DocumentRevision(document_id=document.id, revision=document.revision,
                            content_json=payload, created_by=user.email))
    db.commit()
    result = document_view(document)
    await hub.publish(project_id, {"type": "document.saved", "slot": slot,
                                   "revision": document.revision, "updatedBy": user.email})
    return result


@app.get("/api/projects/{project_id}/documents/{slot}/history")
def history(project_id: int, slot: str, user: User = Depends(current_user),
            db: Session = Depends(session_scope)):
    del user
    document = find_document(db, project_id, slot)
    values = db.scalars(select(DocumentRevision).where(DocumentRevision.document_id == document.id)
                        .order_by(DocumentRevision.created_at.desc()).limit(2000)).all()
    selected = []
    interval = timedelta(minutes=max(1, cfg.history_interval_minutes))
    for item in values:
        item_time = item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at
        previous_time = None if not selected else (selected[-1].created_at.replace(tzinfo=timezone.utc)
                                                    if selected[-1].created_at.tzinfo is None else selected[-1].created_at)
        if previous_time is None or previous_time - item_time >= interval:
            selected.append(item)
        if len(selected) >= max(1, cfg.history_limit):
            break
    return [{"revision": item.revision, "createdAt": iso(item.created_at), "createdBy": item.created_by}
            for item in selected]


@app.get("/api/projects/{project_id}/documents/{slot}/history/{revision}")
def history_document(project_id: int, slot: str, revision: int,
                     user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    document = find_document(db, project_id, slot)
    value = db.scalar(select(DocumentRevision).where(DocumentRevision.document_id == document.id,
                                                       DocumentRevision.revision == revision))
    if not value:
        raise HTTPException(404, "Версия не найдена")
    content = json.loads(value.content_json)
    return {"revision": value.revision, "createdAt": iso(value.created_at),
            "createdBy": value.created_by, "pages": content.get("pages") or {}}


@app.post("/api/projects/{project_id}/documents/{slot}/history/{revision}/restore")
async def restore_document(project_id: int, slot: str, revision: int, body: RestoreBody,
                           user: User = Depends(current_user), db: Session = Depends(session_scope)):
    document = find_document(db, project_id, slot)
    if body.baseRevision != document.revision:
        raise HTTPException(409, "Документ уже изменён другим медиатором")
    source = db.scalar(select(DocumentRevision).where(DocumentRevision.document_id == document.id,
                                                        DocumentRevision.revision == revision))
    if not source:
        raise HTTPException(404, "Версия не найдена")
    document.revision += 1
    document.content_json = source.content_json
    document.updated_at = now()
    document.updated_by = user.email
    db.add(DocumentRevision(document_id=document.id, revision=document.revision,
                            content_json=document.content_json, created_by=user.email))
    db.commit()
    result = document_view(document)
    await hub.publish(project_id, {"type": "document.saved", "slot": slot,
                                   "revision": document.revision, "updatedBy": user.email})
    return result


@app.post("/api/projects/{project_id}/assets")
async def upload_asset(project_id: int, file: UploadFile = File(...),
                       user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    if not db.get(Project, project_id):
        raise HTTPException(404, "Проект не найден")
    extension = Path(file.filename or "").suffix.lower()
    if extension not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(400, "Поддерживаются PNG, JPG и WEBP")
    data = await file.read(25 * 1024 * 1024 + 1)
    if len(data) > 25 * 1024 * 1024:
        raise HTTPException(413, "Изображение больше 25 МБ")
    folder = assets_root / str(project_id)
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{uuid.uuid4().hex}{extension}"
    (folder / name).write_bytes(data)
    return {"url": f"/assets/{project_id}/{name}", "name": name, "size": len(data)}


@app.post("/api/projects/{project_id}/assets/import-docx")
async def import_docx_assets(project_id: int, file: UploadFile = File(...),
                             user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    if not db.get(Project, project_id):
        raise HTTPException(404, "Проект не найден")
    if Path(file.filename or "").suffix.lower() != ".docx":
        raise HTTPException(400, "Выберите таблицу в формате DOCX")
    data = await file.read(60 * 1024 * 1024 + 1)
    if len(data) > 60 * 1024 * 1024:
        raise HTTPException(413, "DOCX больше 60 МБ")
    try:
        source = WordDocument(BytesIO(data))
    except Exception as exc:
        raise HTTPException(400, "Не удалось прочитать DOCX") from exc
    if not source.tables:
        raise HTTPException(400, "В DOCX нет таблиц")
    table = source.tables[-1]
    ordered = []
    occurrences = Counter()
    cells = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            relation_ids = [node.get(qn("r:embed")) for node in cell._tc.xpath(".//a:blip")]
            occurrences.update(relation_ids)
            cells.append((row_index, column_index, " ".join(cell.text.split()), relation_ids))
    seen = set()
    for row_index, column_index, label, relation_ids in cells:
        for relation_id in relation_ids:
            # Repeated image relationships in supplied tables are blank cell placeholders.
            if not relation_id or occurrences[relation_id] > 1 or relation_id in seen:
                continue
            seen.add(relation_id)
            part = source.part.related_parts.get(relation_id)
            content_type = getattr(part, "content_type", "")
            extension = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}.get(content_type)
            if not extension:
                continue
            ordered.append((row_index, column_index, label, part.blob, extension))
    if not ordered:
        raise HTTPException(400, "В последней таблице не найдено изображений")
    folder = assets_root / str(project_id)
    folder.mkdir(parents=True, exist_ok=True)
    result = []
    for number, (row_index, column_index, label, blob, extension) in enumerate(ordered, 1):
        name = f"table-{number:03d}-{uuid.uuid4().hex}{extension}"
        (folder / name).write_bytes(blob)
        result.append({"number": number, "row": row_index + 1, "column": column_index + 1,
                       "label": label, "url": f"/assets/{project_id}/{name}"})
    return {"tableIndex": len(source.tables), "count": len(result), "assets": result}


@app.get("/api/projects/{project_id}/documents/{slot}/template")
def get_document_template(project_id: int, slot: str, user: User = Depends(current_user),
                          db: Session = Depends(session_scope)):
    del user
    document = find_document(db, project_id, slot)
    value = db.scalar(select(DocumentTemplate).where(DocumentTemplate.document_id == document.id))
    if not value:
        raise HTTPException(404, "Сохранённый шаблон ещё не создан")
    content = json.loads(value.content_json)
    return {"pages": content.get("pages") or {}, "updatedAt": iso(value.updated_at),
            "updatedBy": value.updated_by}


@app.put("/api/projects/{project_id}/documents/{slot}/template")
def save_document_template(project_id: int, slot: str, body: SaveTemplateBody,
                           user: User = Depends(publisher), db: Session = Depends(session_scope)):
    document = find_document(db, project_id, slot)
    pages = clean_pages(body.pages)
    payload = json.dumps({"pages": pages}, ensure_ascii=False, separators=(",", ":"))
    if len(payload.encode("utf-8")) > 4 * 1024 * 1024:
        raise HTTPException(413, "Шаблон слишком большой")
    value = db.scalar(select(DocumentTemplate).where(DocumentTemplate.document_id == document.id))
    if value:
        value.content_json = payload
        value.updated_at = now()
        value.updated_by = user.email
    else:
        value = DocumentTemplate(document_id=document.id, content_json=payload,
                                 updated_at=now(), updated_by=user.email)
        db.add(value)
    db.commit()
    return {"saved": True, "updatedAt": iso(value.updated_at), "updatedBy": value.updated_by}


@app.get("/api/template")
def template(user: User = Depends(current_user)):
    del user
    if not cfg.template_pdf.is_file():
        raise HTTPException(503, f"Шаблон не найден: {cfg.template_pdf}")
    return FileResponse(cfg.template_pdf, media_type="application/pdf", filename="Postcards-template.pdf")


@app.get("/api/projects/{project_id}/documents/{slot}/export.pdf")
def export_document(project_id: int, slot: str, user: User = Depends(current_user),
                    db: Session = Depends(session_scope)):
    del user
    document = find_document(db, project_id, slot)
    export_dir = cfg.storage_dir / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    target = export_dir / f"{uuid.uuid4().hex}.pdf"
    render_document(cfg.template_pdf, json.loads(document.content_json), assets_root, target)
    filename = f"{slot}-{document.revision}.pdf"
    return FileResponse(target, media_type="application/pdf", filename=filename,
                        background=BackgroundTask(target.unlink, missing_ok=True))


@app.post("/api/projects/{project_id}/publish")
async def publish(project_id: int, user: User = Depends(publisher), db: Session = Depends(session_scope)):
    project_value = db.get(Project, project_id)
    if not project_value:
        raise HTTPException(404, "Проект не найден")
    documents = db.scalars(select(Document).where(Document.project_id == project_id).order_by(Document.slot)).all()
    if len(documents) != 4:
        raise HTTPException(409, "Для публикации нужны четыре документа")
    sequence = (db.scalar(select(func.max(Release.sequence)).where(Release.project_id == project_id)) or 0) + 1
    stamp = now()
    release_id = f"{stamp.strftime('%Y%m%dT%H%M%SZ')}-r{sequence:04d}-{uuid.uuid4().hex[:8]}"
    doc_values = [document_view(item) for item in documents]
    manifest = publish_release(cfg.template_pdf, assets_root, releases_root, release_id, doc_values, {
        "schemaVersion": 1, "releaseId": release_id, "sequence": sequence,
        "projectId": project_id, "projectName": project_value.name,
        "createdAt": iso(stamp), "createdBy": user.email,
    })
    db.add(Release(id=release_id, project_id=project_id, sequence=sequence,
                   manifest_json=json.dumps(manifest, ensure_ascii=False), created_at=stamp,
                   created_by=user.email))
    db.commit()
    await hub.publish(project_id, {"type": "release.published", "release": manifest})
    return manifest


def release_view(release: Release) -> dict:
    return json.loads(release.manifest_json)


@app.get("/api/projects/{project_id}/releases")
def releases(project_id: int, user: User = Depends(current_user), db: Session = Depends(session_scope)):
    del user
    return [release_view(item) for item in db.scalars(select(Release).where(Release.project_id == project_id)
                                                      .order_by(Release.sequence.desc()).limit(50)).all()]


@app.get("/api/sync/latest")
def sync_latest(access: None = Depends(sync_access), db: Session = Depends(session_scope)):
    del access
    release = db.scalar(select(Release).order_by(Release.created_at.desc()))
    if not release:
        raise HTTPException(404, "Опубликованных релизов пока нет")
    return release_view(release)


@app.get("/api/sync/releases/{release_id}/{filename}")
def sync_file(release_id: str, filename: str, access: None = Depends(sync_access)):
    del access
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", filename):
        raise HTTPException(400, "Недопустимое имя файла")
    path = releases_root / release_id / filename
    if not path.is_file() or releases_root.resolve() not in path.resolve().parents:
        raise HTTPException(404, "Файл релиза не найден")
    return FileResponse(path, media_type="application/pdf" if path.suffix == ".pdf" else "application/json")


@app.websocket("/ws/projects/{project_id}")
async def project_socket(socket: WebSocket, project_id: int, token: str = ""):
    try:
        decode_token(token)
    except HTTPException:
        await socket.close(code=4401)
        return
    await hub.connect(project_id, socket)
    try:
        await socket.send_json({"type": "connected", "projectId": project_id})
        while True:
            await socket.receive_text()
    except WebSocketDisconnect:
        hub.disconnect(project_id, socket)
