from __future__ import annotations

import json
import io
from collections import Counter
from pathlib import Path

import pymupdf as fitz
from docx import Document as WordDocument
from docx.oxml.ns import qn
from PIL import Image, ImageChops, ImageOps, ImageStat, UnidentifiedImageError


PAGE_COUNT = 28
PARAGRAPH_INDENT = "\u2003"
DOCUMENT_TITLES = {
    "pdf1": "Экфрасисы",
    "pdf2": "Формальный анализ",
    "pdf3": "Контекст создания",
    "pdf4": "Контекст восприятия",
}

# Порядок разворотов подтверждён заказчиком 15.10.2026. ``None`` означает
# намеренно чёрную заглушку: оборот в архиве не сохранился.
SPREAD_SCANS = (
    ((1, "1"), (None, "1.1")),
    ((17, "1.21"), (9, "1.2")),
    ((2, "2"), (None, "2.1")),
    ((18, "2.21"), (10, "2.2")),
    ((3, "3"), (11, "3.1")),
    ((4, "4"), (None, "4.1")),
    ((None, "4.21"), (12, "4.2")),
    ((5, "5"), (13, "5.1")),
    ((6, "6"), (None, "6.1")),
    ((19, "6.21"), (14, "6.2")),
    ((7, "7"), (None, "7.1")),
    ((20, "7.21"), (15, "7.2")),
    ((8, "8"), (None, "8.1")),
    ((21, "8.21"), (16, "8.2")),
)
CUSTOMER_SCAN_CELLS = {
    (0, 0): 1, (0, 2): 2, (0, 4): 3, (0, 5): 11, (0, 6): 4,
    (1, 0): 17, (1, 1): 9, (1, 2): 18, (1, 3): 10, (1, 7): 12,
    (2, 0): 5, (2, 1): 13, (2, 2): 6, (2, 4): 7, (2, 6): 8,
    (3, 2): 19, (3, 3): 14, (3, 4): 20, (3, 5): 15, (3, 6): 21, (3, 7): 16,
}
# Соответствие старого 28-страничного pdf2.pdf утверждённым номерам.
# Скан 20 отсутствовал в старом PDF; скан 21, наоборот, берётся из него,
# потому что в DOCX для «Греков, Тачанка» осталась только подпись.
LEGACY_CARD_BY_SCAN = {
    17: 1, 9: 2, 1: 3, 5: 5, 13: 6, 2: 7, 18: 9, 10: 10,
    6: 11, 19: 13, 14: 14, 3: 15, 11: 16, 7: 17, 21: 19,
    15: 20, 4: 21, 12: 24, 16: 26, 8: 27,
}


def _item(item_id: str, item_type: str, x: float, y: float, w: float, h: float, **extra) -> dict:
    return {"id": item_id, "type": item_type, "x": x, "y": y, "w": w, "h": h, **extra}


def blank_editable_layout(page_count: int = PAGE_COUNT) -> dict[str, list[dict]]:
    """Create the 14-spread layout used by the installation.

    Odd pages contain the scans, captions and the beginning of each text. Even
    pages are continuation pages. The geometry is locked, while text and image
    contents remain editable in the browser.
    """
    pages: dict[str, list[dict]] = {}
    for page in range(1, page_count + 1):
        items = [_item(f"background-{page}", "page-background", 0, 0, 1, 1,
                       color="#ffffff", locked=True)]
        for column in range(2):
            card_index = ((page - 1) // 2) * 2 + column + 1
            scan_number, scan_code = SPREAD_SCANS[(page - 1) // 2][column]
            # Both page halves meet exactly at the centre, as in the approved
            # Illustrator sample. This avoids two parallel contours and makes
            # the scan/caption/body sections one continuous column grid.
            x = 0.018 if column == 0 else 0.500
            column_width = 0.482
            inner_width = column_width - 0.010
            items.append(_item(
                f"outline-{page}-{column}", "shape", x, 0.032, column_width, 0.936,
                shape="rectangle", fill="transparent", stroke="#777777", strokeWidth=1,
                strokeStyle="dashed", locked=True, role="outline", cardIndex=card_index,
            ))
            if page % 2:
                items.extend([
                    _item(
                        f"scan-{page}-{column}", "image", x + 0.062, 0.034, 0.358, 0.337,
                        src="", fit="contain", locked=True, role="scan", sequence=scan_number,
                        scanCode=scan_code, missing=scan_number is None, cardIndex=card_index,
                    ),
                    _item(
                        f"scan-divider-{page}-{column}", "shape", x + 0.005, 0.382, inner_width, 0.006,
                        shape="line", fill="transparent", stroke="#777777", strokeWidth=1,
                        strokeStyle="dashed", locked=True, role="scan-divider", cardIndex=card_index,
                    ),
                    _item(
                        f"caption-{page}-{column}", "text", x + 0.005, 0.393, inner_width, 0.078,
                        text="", placeholder="", fontSize=16,
                        fontFamily="Times New Roman", color="#111111", background="transparent",
                        align="right", bold=False, italic=False, locked=True, role="caption",
                        cardIndex=card_index,
                    ),
                    _item(
                        f"caption-divider-{page}-{column}", "shape", x + 0.005, 0.476, inner_width, 0.006,
                        shape="line", fill="transparent", stroke="#777777", strokeWidth=1,
                        strokeStyle="dashed", locked=True, role="caption-divider", cardIndex=card_index,
                    ),
                    _item(
                        f"body-{page}-{column}", "text", x + 0.005, 0.489, inner_width, 0.473,
                        text="", placeholder="", fontSize=15, fontFamily="Times New Roman",
                        color="#111111", background="transparent", align="left", bold=False,
                        italic=False, locked=True, role="body", cardIndex=card_index,
                    ),
                ])
            else:
                items.append(_item(
                    f"body-{page}-{column}", "text", x + 0.005, 0.025, inner_width, 0.935,
                    text="", placeholder="", fontSize=15,
                    fontFamily="Times New Roman", color="#111111", background="transparent",
                    align="left", bold=False, italic=False, locked=True, role="body",
                    cardIndex=card_index,
                ))
        pages[str(page)] = items
    return pages


def indent_paragraphs(text: str) -> str:
    """Prefix every non-empty paragraph with one em-space (a red line)."""
    return "\n".join(
        line if not line.strip() or line.startswith(PARAGRAPH_INDENT) else PARAGRAPH_INDENT + line.lstrip()
        for line in str(text or "").split("\n")
    )


def _region_text(page: fitz.Page, x0: float, x1: float, y0: float, y1: float,
                 reflow: bool = False) -> str:
    rows: list[list] = []
    for trace in page.get_texttrace():
        for codepoint, _glyph, origin, bbox in trace["chars"]:
            x, y = origin
            if not (x0 <= x < x1 and y0 <= y < y1 and bbox[1] < page.rect.height and bbox[3] > 0):
                continue
            if codepoint < 0 or codepoint > 0x10FFFF:
                continue
            row = next((value for value in rows if abs(value[0] - y) < 1), None)
            if row is None:
                row = [y, []]
                rows.append(row)
            row[1].append((x, chr(codepoint)))
    lines: list[tuple[float, float, str]] = []
    for _y, chars in sorted(rows, key=lambda value: value[0]):
        chars = sorted(chars, key=lambda value: value[0])
        line = "".join(char for _x, char in chars).strip()
        if line:
            lines.append((_y, chars[0][0], line))
    if not reflow:
        return "\n".join(line for _y, _x, line in lines)
    paragraphs: list[str] = []
    base_x = min((line_x for _line_y, line_x, _line in lines), default=x0)
    for index, (line_y, line_x, line) in enumerate(lines):
        previous_y = lines[index - 1][0] if index else line_y
        begins_paragraph = index == 0 or line_x > base_x + 8 or line_y - previous_y > 24
        if begins_paragraph:
            paragraphs.append(line)
        elif paragraphs[-1].endswith("-"):
            paragraphs[-1] = paragraphs[-1][:-1] + line
        else:
            paragraphs[-1] += " " + line
    return indent_paragraphs("\n".join(paragraphs))


def import_legacy_pdf(source_pdf: Path, assets_root: Path, project_id: int = 1) -> dict[str, list[dict]]:
    """Convert the legacy 28-page PDF into editable page objects and scan PNGs."""
    source_pdf = Path(source_pdf)
    document = fitz.open(source_pdf)
    if len(document) != PAGE_COUNT:
        document.close()
        raise ValueError(f"Ожидалось {PAGE_COUNT} страниц, найдено {len(document)}")

    pages = blank_editable_layout()
    asset_folder = Path(assets_root) / str(project_id) / "starter-v1"
    asset_folder.mkdir(parents=True, exist_ok=True)
    width = document[0].rect.width

    for page_index, page in enumerate(document):
        page_number = page_index + 1
        for column in range(2):
            x0 = column * width / 2
            x1 = (column + 1) * width / 2
            card_index = (page_index // 2) * 2 + column + 1
            body = _region_text(page, x0, x1, 405 if page_number % 2 else 0,
                                page.rect.height, reflow=True)
            body_item = next(item for item in pages[str(page_number)]
                             if item.get("role") == "body" and item.get("cardIndex") == card_index)
            body_item["text"] = body

            if page_number % 2 == 0:
                continue
            caption = _region_text(page, x0, x1, 315, 405)
            caption_item = next(item for item in pages[str(page_number)]
                                if item.get("role") == "caption" and item.get("cardIndex") == card_index)
            caption_item["text"] = caption

            # Real scans are placed by their catalogue number. Missing reverse
            # sides remain generated black placeholders instead of bitmap files.
            image_item = next(item for item in pages[str(page_number)]
                              if item.get("role") == "scan" and item.get("cardIndex") == card_index)
            clip = fitz.Rect(x0 + 90, 27, x0 + 518, 313)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), clip=clip, alpha=False)
            pixmap.save(asset_folder / f"legacy-card-{card_index:02d}.png")
            if image_item.get("missing"):
                continue
            name = f"scan-{int(image_item['sequence']):02d}.png"
            pixmap.save(asset_folder / name)
            image_item["src"] = f"/assets/{project_id}/starter-v1/{name}"
            image_item["sourceLabel"] = caption.replace("\n", " ")
    document.close()
    return pages


def _decoded_image(blob: bytes) -> Image.Image | None:
    try:
        with Image.open(io.BytesIO(blob)) as image:
            return image.convert("RGB")
    except (UnidentifiedImageError, OSError):
        return None


def _visual_score(left: Image.Image, right: Image.Image) -> float:
    left_small = ImageOps.fit(left, (64, 64))
    right_small = ImageOps.fit(right, (64, 64))
    return sum(value * value for value in ImageStat.Stat(ImageChops.difference(left_small, right_small)).rms)


def _best_original(part, candidates: list) -> object:
    """Resolve a Word table thumbnail to the larger copy embedded elsewhere in the DOCX."""
    thumbnail = _decoded_image(part.blob)
    if thumbnail is None:
        return part
    thumbnail_area = thumbnail.width * thumbnail.height
    matches = []
    for candidate in candidates:
        image = _decoded_image(candidate.blob)
        if image is None or image.width * image.height < thumbnail_area * 4:
            continue
        matches.append((_visual_score(thumbnail, image), -(image.width * image.height), candidate))
    if not matches:
        return part
    score, _area, match = min(matches, key=lambda value: (value[0], value[1]))
    # Resized copies of the same scan score in the low hundreds. A conservative
    # threshold prevents visually similar blank postcard backs from being mixed up.
    return match if score <= 750 else part


def extract_customer_scans(source_docx: Path, assets_root: Path, project_id: int = 1,
                           asset_version: str = "starter-v3") -> dict[int, dict]:
    source = WordDocument(source_docx)
    if not source.tables:
        raise ValueError("В DOCX нет таблиц")
    table = source.tables[-1]
    if len(table.rows) != 4 or len(table.columns) != 8:
        raise ValueError("Ожидалась утверждённая таблица 4×8")
    occurrences: Counter[str] = Counter()
    cells = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            relation_ids = [node.get(qn("r:embed")) for node in cell._tc.xpath(".//a:blip")]
            occurrences.update(relation_ids)
            cells.append((row_index, column_index, " ".join(cell.text.split()), relation_ids))
    image_parts = []
    known_part_names = set()
    for part in source.part.related_parts.values():
        content_type = getattr(part, "content_type", "")
        part_name = str(getattr(part, "partname", ""))
        if content_type.startswith("image/") and part_name not in known_part_names:
            known_part_names.add(part_name)
            image_parts.append(part)
    folder = Path(assets_root) / str(project_id) / asset_version
    folder.mkdir(parents=True, exist_ok=True)
    result: dict[int, dict] = {}
    seen = set()
    for row_index, column_index, label, relation_ids in cells:
        number = CUSTOMER_SCAN_CELLS.get((row_index, column_index))
        if not number:
            continue
        for relation_id in relation_ids:
            if not relation_id or occurrences[relation_id] > 1 or relation_id in seen:
                continue
            seen.add(relation_id)
            part = source.part.related_parts.get(relation_id)
            part = _best_original(part, image_parts)
            content_type = getattr(part, "content_type", "")
            extension = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}.get(content_type)
            if not extension:
                continue
            name = f"scan-{number:02d}{extension}"
            path = folder / name
            path.write_bytes(part.blob)
            result[number] = {"path": path, "url": f"/assets/{project_id}/{asset_version}/{name}", "label": label}
            break
    required = set(range(1, 21))
    if not required.issubset(result):
        raise ValueError(f"В таблице найдены не все обязательные сканы 1–20: {sorted(result)}")
    return result


def prefer_legacy_for_small_scans(scans: dict[int, dict], assets_root: Path, project_id: int = 1,
                                  asset_version: str = "starter-v3") -> None:
    """Use the sharper legacy-PDF rendering when the DOCX only contains a tiny preview."""
    legacy_folder = Path(assets_root) / str(project_id) / "starter-v1"
    target_folder = Path(assets_root) / str(project_id) / asset_version
    target_folder.mkdir(parents=True, exist_ok=True)
    for number, value in list(scans.items()):
        old_card = LEGACY_CARD_BY_SCAN.get(number)
        if not old_card:
            continue
        image = _decoded_image(value["path"].read_bytes())
        if image is not None and image.width >= 400 and image.height >= 300:
            continue
        source = legacy_folder / f"legacy-card-{old_card:02d}.png"
        if not source.is_file():
            continue
        target = target_folder / f"scan-{number:02d}.png"
        target.write_bytes(source.read_bytes())
        value.update(path=target, url=f"/assets/{project_id}/{asset_version}/{target.name}")


def import_customer_layout(source_pdf: Path, source_docx: Path, assets_root: Path,
                           project_id: int = 1) -> dict[str, list[dict]]:
    """Rebuild the editable layout, preserving legacy text and placing scans by approved numbering."""
    legacy = import_legacy_pdf(source_pdf, assets_root, project_id)
    scans = extract_customer_scans(source_docx, assets_root, project_id)
    legacy_folder = Path(assets_root) / str(project_id) / "starter-v1"
    prefer_legacy_for_small_scans(scans, assets_root, project_id)
    if 21 not in scans:
        source = legacy_folder / f"legacy-card-{LEGACY_CARD_BY_SCAN[21]:02d}.png"
        target_folder = Path(assets_root) / str(project_id) / "starter-v3"
        target = target_folder / "scan-21.png"
        target.write_bytes(source.read_bytes())
        scans[21] = {"path": target, "url": f"/assets/{project_id}/starter-v3/scan-21.png",
                     "label": "А. М. Греков. Тачанка"}
    pages = blank_editable_layout()

    def copy_text(old_card: int, new_card: int) -> None:
        old_spread = (old_card - 1) // 2
        new_spread = (new_card - 1) // 2
        for old_page, new_page in ((old_spread * 2 + 1, new_spread * 2 + 1),
                                   (old_spread * 2 + 2, new_spread * 2 + 2)):
            for role in ("caption", "body"):
                old_item = next((item for item in legacy[str(old_page)]
                                 if item.get("role") == role and item.get("cardIndex") == old_card), None)
                new_item = next((item for item in pages[str(new_page)]
                                 if item.get("role") == role and item.get("cardIndex") == new_card), None)
                if old_item and new_item:
                    new_item["text"] = old_item.get("text", "")

    for values in pages.values():
        for item in values:
            if item.get("role") != "scan" or item.get("missing"):
                continue
            number = int(item["sequence"])
            if number not in scans:
                continue
            item["src"] = scans[number]["url"]
            item["sourceLabel"] = scans[number]["label"]
            old_card = LEGACY_CARD_BY_SCAN.get(number)
            if old_card:
                copy_text(old_card, int(item["cardIndex"]))
                old_pair = old_card + 1 if old_card % 2 else old_card - 1
                target_pair = int(item["cardIndex"]) + 1 if int(item["cardIndex"]) % 2 else int(item["cardIndex"]) - 1
                target_pair_item = next((candidate for candidate in values
                                         if candidate.get("role") == "scan" and candidate.get("cardIndex") == target_pair), None)
                if target_pair_item and target_pair_item.get("missing"):
                    copy_text(old_pair, target_pair)
    return pages


def without_body_text(pages: dict[str, list[dict]]) -> dict[str, list[dict]]:
    clone = json.loads(json.dumps(pages, ensure_ascii=False))
    for items in clone.values():
        for item in items:
            if item.get("role") == "body":
                item["text"] = ""
    return clone
