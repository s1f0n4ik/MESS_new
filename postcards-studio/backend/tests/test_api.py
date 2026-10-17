import hashlib
import importlib
import json
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import pymupdf as fitz
from docx import Document as WordDocument
from docx.shared import Inches
from fastapi.testclient import TestClient


class StudioApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        root = Path(cls.temp.name)
        template = root / "template.pdf"
        pdf = fitz.open()
        for _ in range(28):
            pdf.new_page(width=1190.55, height=841.89)
        pdf[0].insert_text((50, 50), "ORIGINAL TEMPLATE")
        pdf.save(template)
        pdf.close()
        os.environ.update({
            "STUDIO_STORAGE_DIR": str(root / "data"),
            "STUDIO_DATABASE_URL": f"sqlite:///{(root / 'studio.db').as_posix()}",
            "STUDIO_TEMPLATE_PDF": str(template),
            "STUDIO_SECRET_KEY": "test-secret-longer-than-thirty-two-bytes",
            "STUDIO_ADMIN_EMAIL": "admin@test.local",
            "STUDIO_ADMIN_PASSWORD": "very-secure-test-password",
            "STUDIO_SYNC_TOKEN": "sync-test-token",
        })
        from app import config
        config.settings.cache_clear()
        from app import main
        importlib.reload(main)
        cls.client = TestClient(main.app)
        cls.client.__enter__()
        login = cls.client.post("/api/auth/login", json={"email": "admin@test.local",
                                                          "password": "very-secure-test-password"})
        assert login.status_code == 200, login.text
        cls.token = login.json()["token"]
        cls.headers = {"Authorization": f"Bearer {cls.token}"}

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        from app.db import engine
        engine.dispose()
        cls.temp.cleanup()

    def test_document_conflict_and_publication(self):
        project = self.client.get("/api/projects/1", headers=self.headers)
        self.assertEqual(project.status_code, 200)
        self.assertEqual(len(project.json()["documents"]), 4)
        body = {"baseRevision": 0, "pages": {"1": [{"id": "hello", "type": "text",
                 "x": .1, "y": .1, "w": .4, "h": .2, "text": "Привет",
                 "fontSize": 28, "color": "#111111", "background": "transparent"}]}}
        saved = self.client.put("/api/projects/1/documents/pdf1", headers=self.headers, json=body)
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json()["revision"], 1)
        conflict = self.client.put("/api/projects/1/documents/pdf1", headers=self.headers, json=body)
        self.assertEqual(conflict.status_code, 409)
        history = self.client.get("/api/projects/1/documents/pdf1/history", headers=self.headers)
        self.assertEqual(history.json()[0]["revision"], 1)
        restored = self.client.post("/api/projects/1/documents/pdf1/history/1/restore",
                                    headers=self.headers, json={"baseRevision": 1})
        self.assertEqual(restored.status_code, 200)
        self.assertEqual(restored.json()["revision"], 2)
        created = self.client.post("/api/users", headers=self.headers, json={
            "email": "mediator@test.local", "name": "Медиатор", "password": "mediator-password",
            "role": "editor",
        })
        self.assertEqual(created.status_code, 200)
        self.assertEqual(len(self.client.get("/api/users", headers=self.headers).json()), 2)
        mediator_login = self.client.post("/api/auth/login", json={
            "email": "mediator@test.local", "password": "mediator-password",
        })
        self.assertEqual(mediator_login.status_code, 200)
        mediator_headers = {"Authorization": f"Bearer {mediator_login.json()['token']}"}
        renamed = self.client.patch("/api/projects/1/documents/pdf1/title", headers=self.headers,
                                    json={"title": "  Экфрасисы выставки  "})
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["title"], "Экфрасисы выставки")
        project_after_rename = self.client.get("/api/projects/1", headers=self.headers).json()
        self.assertEqual(project_after_rename["documents"][0]["title"], "Экфрасисы выставки")
        forbidden_rename = self.client.patch("/api/projects/1/documents/pdf1/title",
                                             headers=mediator_headers, json={"title": "Нельзя"})
        self.assertEqual(forbidden_rename.status_code, 403)
        forbidden = self.client.delete(f"/api/users/{created.json()['id']}", headers=mediator_headers)
        self.assertEqual(forbidden.status_code, 403)
        current_id = self.client.get("/api/me", headers=self.headers).json()["id"]
        self_delete = self.client.delete(f"/api/users/{current_id}", headers=self.headers)
        self.assertEqual(self_delete.status_code, 409)
        service_headers = {"X-Collaboration-Secret": "development-collaboration-secret"}
        collab_auth = self.client.post("/api/internal/collaboration/auth", headers=service_headers,
                                       json={"token": self.token, "documentName": "project-1-pdf1-layout-v5"})
        self.assertEqual(collab_auth.status_code, 200, collab_auth.text)
        self.assertEqual(collab_auth.json()["slot"], "pdf1")
        stale_collab_auth = self.client.post("/api/internal/collaboration/auth", headers=service_headers,
                                             json={"token": self.token, "documentName": "project-1-pdf1"})
        self.assertEqual(stale_collab_auth.status_code, 400)
        collab_snapshot = self.client.post("/api/internal/collaboration/documents/project-1-pdf1-layout-v5",
                                           headers=service_headers,
                                           json={"pages": body["pages"], "updatedBy": "admin@test.local"})
        self.assertEqual(collab_snapshot.status_code, 200, collab_snapshot.text)
        unauthorized = self.client.get("/api/internal/collaboration/documents/project-1-pdf1-layout-v5")
        self.assertEqual(unauthorized.status_code, 401)
        deleted = self.client.delete(f"/api/users/{created.json()['id']}", headers=self.headers)
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(deleted.json()["deleted"], created.json()["id"])
        self.assertEqual(len(self.client.get("/api/users", headers=self.headers).json()), 1)
        revoked = self.client.get("/api/me", headers=mediator_headers)
        self.assertEqual(revoked.status_code, 401)
        release = self.client.post("/api/projects/1/publish", headers=self.headers)
        self.assertEqual(release.status_code, 200, release.text)
        self.assertEqual(len(release.json()["files"]), 4)
        latest = self.client.get("/api/sync/latest", headers={"X-Sync-Token": "sync-test-token"})
        self.assertEqual(latest.json()["releaseId"], release.json()["releaseId"])
        for item in release.json()["files"]:
            response = self.client.get(f"/api/sync/releases/{release.json()['releaseId']}/{item['name']}",
                                       headers={"X-Sync-Token": "sync-test-token"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(hashlib.sha256(response.content).hexdigest(), item["sha256"])

    def test_extended_editor_elements_are_sanitized_and_rendered(self):
        from app.main import clean_pages
        from app.pdf_renderer import render_document

        pages = clean_pages({"1": [
            {"id": "background-1", "type": "page-background", "x": 0, "y": 0, "w": 1, "h": 1,
             "color": "#fefefe"},
            {"id": "styled-text", "type": "text", "x": .1, "y": .1, "w": .6, "h": .2,
             "text": "Курсивный текст", "fontSize": 24, "fontFamily": "Georgia", "italic": True,
             "bold": True, "color": "#223344", "background": "transparent", "align": "center"},
            {"id": "dashed-shape", "type": "shape", "x": .15, "y": .4, "w": .4, "h": .25,
             "shape": "ellipse", "fill": "transparent", "stroke": "#112233", "strokeWidth": 3,
             "strokeStyle": "dashed"},
        ]})
        self.assertEqual(pages["1"][0]["type"], "page-background")
        self.assertEqual(pages["1"][1]["fontFamily"], "Georgia")
        self.assertTrue(pages["1"][1]["italic"])
        self.assertEqual(pages["1"][2]["strokeStyle"], "dashed")

        root = Path(self.temp.name)
        output = root / "extended.pdf"
        render_document(root / "template.pdf", {"pages": pages}, root / "data" / "assets", output)
        self.assertTrue(output.is_file())
        rendered = fitz.open(output)
        self.assertGreater(len(rendered[0].get_drawings()), 0)
        self.assertNotIn("ORIGINAL TEMPLATE", rendered[0].get_text())
        self.assertIn("Курсивный текст", rendered[0].get_text().replace("\xa0", " "))
        self.assertFalse(any(font[3] == "Helvetica" for font in rendered[0].get_fonts(full=True)))
        rendered.close()

    def test_collaboration_history_is_throttled_and_document_can_be_exported(self):
        first_pages = {"1": [{"id": "first", "type": "text", "x": .1, "y": .1,
                               "w": .4, "h": .2, "text": "Первая версия", "fontSize": 24,
                               "color": "#111111", "background": "transparent"}]}
        second_pages = {"1": [{**first_pages["1"][0], "text": "Вторая версия"}]}
        first = self.client.post("/api/projects/1/documents/pdf2/collaboration-snapshot",
                                 headers=self.headers, json={"pages": first_pages})
        second = self.client.post("/api/projects/1/documents/pdf2/collaboration-snapshot",
                                  headers=self.headers, json={"pages": second_pages})
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertGreater(second.json()["revision"], first.json()["revision"])
        history = self.client.get("/api/projects/1/documents/pdf2/history", headers=self.headers)
        self.assertEqual(len(history.json()), 1)
        checkpoint = self.client.post("/api/projects/1/documents/pdf2/collaboration-snapshot",
                                      headers=self.headers,
                                      json={"pages": second_pages, "checkpoint": True})
        self.assertEqual(checkpoint.status_code, 200)
        history = self.client.get("/api/projects/1/documents/pdf2/history", headers=self.headers)
        self.assertEqual(history.json()[0]["revision"], checkpoint.json()["revision"])

        exported = self.client.get("/api/projects/1/documents/pdf2/export.pdf", headers=self.headers)
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertEqual(exported.headers["content-type"], "application/pdf")
        result = fitz.open(stream=exported.content, filetype="pdf")
        self.assertEqual(len(result), 28)
        self.assertIn("Вторая версия", result[0].get_text().replace("\xa0", " "))
        result.close()

    def test_template_and_word_table_import(self):
        template_pages = {"1": [
            {"id": "background-1", "type": "page-background", "x": 0, "y": 0, "w": 1, "h": 1,
             "color": "#ffffff", "locked": True},
            {"id": "scan-1-0", "type": "image", "x": .05, "y": .05, "w": .4, "h": .3,
             "src": "", "fit": "contain", "locked": True, "role": "scan", "sequence": 1,
             "cardIndex": 1},
            {"id": "body-1-0", "type": "text", "x": .05, "y": .4, "w": .4, "h": .5,
             "text": "Редактируемое поле", "fontSize": 18, "fontFamily": "Times New Roman",
             "color": "#111111", "background": "transparent", "align": "left",
             "locked": True, "role": "body", "cardIndex": 1},
        ]}
        saved = self.client.put("/api/projects/1/documents/pdf1/template", headers=self.headers,
                                json={"pages": template_pages})
        self.assertEqual(saved.status_code, 200, saved.text)
        loaded = self.client.get("/api/projects/1/documents/pdf1/template", headers=self.headers)
        self.assertEqual(loaded.status_code, 200, loaded.text)
        self.assertEqual(loaded.json()["pages"]["1"][1]["role"], "scan")
        self.assertEqual(loaded.json()["pages"]["1"][1]["src"], "")
        self.assertEqual(loaded.json()["pages"]["1"][2]["fontFamily"], "Times New Roman")

        root = Path(self.temp.name)
        red = root / "red.png"
        blue = root / "blue.png"
        placeholder = root / "placeholder.png"
        for target, color in ((red, (255, 0, 0)), (blue, (0, 0, 255)), (placeholder, (255, 255, 255))):
            pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
            for y in range(8):
                for x in range(8):
                    pixmap.set_pixel(x, y, color)
            pixmap.save(target)
        word = WordDocument()
        word.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(str(red), width=Inches(.2))
        table = word.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "Первый"
        table.cell(0, 0).paragraphs[0].add_run().add_picture(str(red), width=Inches(.2))
        table.cell(0, 1).paragraphs[0].add_run().add_picture(str(placeholder), width=Inches(.2))
        table.cell(0, 2).text = "Второй"
        table.cell(0, 2).paragraphs[0].add_run().add_picture(str(placeholder), width=Inches(.2))
        table.cell(0, 2).paragraphs[0].add_run().add_picture(str(blue), width=Inches(.2))
        buffer = BytesIO()
        word.save(buffer)
        imported = self.client.post("/api/projects/1/assets/import-docx", headers=self.headers,
                                    files={"file": ("table.docx", buffer.getvalue(),
                                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(imported.status_code, 200, imported.text)
        self.assertEqual(imported.json()["count"], 2)
        self.assertEqual([item["label"] for item in imported.json()["assets"]], ["Первый", "Второй"])
        self.assertEqual([item["number"] for item in imported.json()["assets"]], [1, 2])


if __name__ == "__main__":
    unittest.main()
